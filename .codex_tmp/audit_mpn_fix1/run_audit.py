from __future__ import annotations

import argparse
import json
import os
import re
import statistics
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "audit_da_compilare_MPN.xlsx"
WORK_DIR = ROOT / ".codex_tmp" / "audit_mpn_fix1"
RESULTS = WORK_DIR / "audit_results.jsonl"
SUMMARY = WORK_DIR / "audit_summary.json"
EXPECTED_RELEASE = "PT26_CLIMA_PROD_2_CAPACITY_MAPPING_FIX1"


def load_rows() -> list[dict[str, Any]]:
    workbook = load_workbook(INPUT, read_only=True, data_only=False)
    sheet = workbook.active
    rows = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        title = str(values[0] or "").strip()
        if not title:
            continue
        rows.append({
            "excel_row": excel_row,
            "title": title,
            "original_mpn": values[1] if len(values) > 1 else None,
        })
    workbook.close()
    return rows


def role_of(item: dict[str, Any]) -> str:
    if item.get("role"):
        return str(item["role"]).upper()
    if item.get("is_ui"):
        return "UI"
    if item.get("is_ue"):
        return "UE"
    return "ACCESSORY"


def numeric_pt(item: dict[str, Any]) -> str | None:
    value = str(item.get("code") or item.get("pt") or "")
    return value if re.fullmatch(r"\d{8}", value) else None


def resolution_type(result: dict[str, Any], legacy_fallback: bool) -> str:
    evidence = result.get("master_evidence") or (result.get("query_analysis") or {}).get("master_evidence") or {}
    reason = evidence.get("reason") if isinstance(evidence, dict) else None
    if reason in {"MASTER_EXACT_IDENTITY", "MASTER_EXACT_PAIR", "MASTER_EXACT_CONFIGURATION"}:
        return reason
    status = result.get("configuration_status")
    if status == "VERIFIED_MASTER_PAIR":
        return "MASTER_EXACT_PAIR"
    if status == "VERIFIED_FULL_COMBINATION":
        return "MASTER_EXACT_CONFIGURATION"
    if status == "NOT_APPLICABLE_EXACT_IDENTITY":
        return "MASTER_EXACT_IDENTITY"
    return "LEGACY_CANDIDATE_MASTER_VALIDATED" if (legacy_fallback and result.get("mpn_final_allowed")) else (
        "LEGACY_CANDIDATE_NOT_CONFIRMED" if legacy_fallback else "MASTER_UNRESOLVED"
    )


def process(engine: Any, row: dict[str, Any]) -> dict[str, Any]:
    started = time.perf_counter()
    base = dict(row)
    try:
        result = engine.search(row["title"], limit=20, include_accessories=True)
        bom = result.get("bom") or []
        main = [item for item in bom if role_of(item) in {"UI", "UE"}]
        accessories = [item for item in bom if role_of(item) not in {"UI", "UE"}]
        main_pts = [pt for item in main if (pt := numeric_pt(item))]
        accessory_pts = [pt for item in accessories if (pt := numeric_pt(item))]
        all_pts = main_pts + accessory_pts
        allowed = bool(result.get("mpn_final_allowed"))
        legacy_fallback = result.get("match_type") != "clima_master_authoritative"
        validation = result.get("master_validation") or {}
        evidence = result.get("master_evidence") or (result.get("query_analysis") or {}).get("master_evidence") or {}
        configuration_key = validation.get("configuration_key") or evidence.get("configuration_key")
        identity = result.get("product_identity_status")
        if not main_pts:
            main_status = "NO_MAIN_UNIT_SELECTED"
        else:
            main_status = identity or "DISCOVERY_ONLY"
        return {
            **base,
            "mpn": " + ".join(all_pts) if allowed else "",
            "codici_pt": "+".join(all_pts),
            "codici_pt_main_unit": "+".join(main_pts),
            "codici_pt_accessori": "+".join(accessory_pts),
            "configuration_status": result.get("configuration_status") or "NOT_APPLICABLE",
            "mpn_final_allowed": allowed,
            "master_resolution_type": resolution_type(result, legacy_fallback),
            "dataset_release": result.get("dataset_release"),
            "main_unit_match_status": main_status,
            "accessory_status": result.get("accessory_status") or "NOT_APPLICABLE",
            "legacy_fallback_used": legacy_fallback,
            "master_configuration_key": configuration_key,
            "runtime_ms": round((time.perf_counter() - started) * 1000, 2),
            "error": "",
        }
    except Exception as exc:
        return {
            **base,
            "mpn": "",
            "codici_pt": "",
            "codici_pt_main_unit": "",
            "codici_pt_accessori": "",
            "configuration_status": "ERROR",
            "mpn_final_allowed": False,
            "master_resolution_type": "ERROR",
            "dataset_release": None,
            "main_unit_match_status": "RUNTIME_ERROR",
            "accessory_status": "ERROR",
            "legacy_fallback_used": False,
            "master_configuration_key": None,
            "runtime_ms": round((time.perf_counter() - started) * 1000, 2),
            "error": f"{type(exc).__name__}: {exc}",
        }


def run_shard(index: int, count: int, resume: bool = False, target_per_shard: int | None = None) -> None:
    if os.environ.get("CLIMA_MASTER_RUNTIME_ENABLED", "1").strip().lower() not in {"1", "true", "yes", "on"}:
        raise SystemExit("CLIMA master runtime is disabled")
    rows = load_rows()
    shard = [row for position, row in enumerate(rows) if position % count == index]
    sys.path.insert(0, str(ROOT))
    os.chdir(ROOT)
    from catalog_search_engine import CatalogSearchEngine

    engine = CatalogSearchEngine()
    engine._ensure_initialized()
    resolver = engine._clima_master_resolver
    if resolver is None or resolver.DATASET_VERSION != EXPECTED_RELEASE:
        raise SystemExit("FIX1 resolver is not active")
    output = WORK_DIR / f"part_{index}.jsonl"
    existing_rows = []
    if resume and output.exists():
        with output.open("r", encoding="utf-8") as handle:
            existing_rows = [json.loads(line) for line in handle if line.strip()]
        completed_excel_rows = {row["excel_row"] for row in existing_rows}
        shard = [row for row in shard if row["excel_row"] not in completed_excel_rows]
        if target_per_shard is not None:
            shard = shard[:max(0, target_per_shard - len(existing_rows))]
    mode = "a" if resume else "w"
    with output.open(mode, encoding="utf-8", newline="\n", buffering=1) as handle:
        for completed, row in enumerate(shard, start=1):
            handle.write(json.dumps(process(engine, row), ensure_ascii=False, separators=(",", ":")) + "\n")
            if completed % 25 == 0 or completed == len(shard):
                print(json.dumps({"shard": index, "completed": completed, "total": len(shard)}), flush=True)


def summarize(count: int, limit: int | None = None) -> None:
    rows = []
    for index in range(count):
        path = WORK_DIR / f"part_{index}.jsonl"
        with path.open("r", encoding="utf-8") as handle:
            rows.extend(json.loads(line) for line in handle if line.strip())
    rows.sort(key=lambda row: row["excel_row"])
    source_rows = load_rows()
    if limit is not None:
        rows = rows[:limit]
        source_rows = source_rows[:limit]
    if len(rows) != len(source_rows):
        raise SystemExit(f"Result count mismatch: {len(rows)} != {len(source_rows)}")
    if [row["excel_row"] for row in rows] != [row["excel_row"] for row in source_rows]:
        raise SystemExit("Result row sequence mismatch")
    if any(row.get("dataset_release") not in {EXPECTED_RELEASE, None} for row in rows):
        raise SystemExit("Unexpected dataset release in results")
    with RESULTS.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
    runtimes = [float(row["runtime_ms"]) for row in rows]
    summary = {
        "rows": len(rows),
        "mpn_final_allowed": sum(bool(row["mpn_final_allowed"]) for row in rows),
        "with_candidate_pt": sum(bool(row["codici_pt"]) for row in rows),
        "errors": sum(bool(row["error"]) for row in rows),
        "configuration_status": Counter(row["configuration_status"] for row in rows),
        "master_resolution_type": Counter(row["master_resolution_type"] for row in rows),
        "legacy_fallback_used": sum(bool(row["legacy_fallback_used"]) for row in rows),
        "runtime_ms_total": round(sum(runtimes), 2),
        "runtime_ms_mean": round(statistics.mean(runtimes), 2),
        "dataset_release": EXPECTED_RELEASE,
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--shard", type=int)
    parser.add_argument("--shard-count", type=int, default=4)
    parser.add_argument("--summarize", action="store_true")
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--target-per-shard", type=int)
    parser.add_argument("--limit", type=int)
    args = parser.parse_args()
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    if args.summarize:
        summarize(args.shard_count, args.limit)
    elif args.shard is not None:
        run_shard(args.shard, args.shard_count, args.resume, args.target_per_shard)
    else:
        parser.error("choose --shard or --summarize")


if __name__ == "__main__":
    main()
