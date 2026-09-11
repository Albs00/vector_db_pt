import argparse
import collections
import hashlib
import json
import math
import os
import re
import statistics
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Principale\Documents\ChatGPT\vector_db_pt")
EXCEL = ROOT / "audit_da_compilare_MPN.xlsx"
FROZEN = ROOT / "benchmark_clima_input_frozen.jsonl"
OUTPUT = ROOT / "benchmark_clima_before_v3.jsonl"
SUMMARY = ROOT / "benchmark_clima_before_v3_summary.json"
MANIFEST = ROOT / "benchmark_clima_before_v3_manifest.json"
PROGRESS = ROOT / ".codex_tmp" / "benchmark_before_v3" / "progress.json"


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_dump_line(record):
    return json.dumps(record, ensure_ascii=False, separators=(",", ":"))


def classify_scope(title):
    text = " ".join(str(title or "").upper().split())
    if "SERVIZIO DI INSTALLAZIONE" in text:
        return "NON_CLIMA", "SERVICE_NOT_PRODUCT"
    if re.search(r"\bCLIMATIZZAT(?:ORE|ORI|RICE|RICI)\b", text):
        return "CLIMA", "EXPLICIT_CLIMATIZZATORE"
    if re.search(r"\bCONDIZIONAT(?:ORE|ORI|RICE|RICI|AMENTO)\b", text):
        return "CLIMA", "EXPLICIT_CONDIZIONAMENTO"
    if re.search(r"UNIT.\s+(?:INTERNA|ESTERNA)", text) and re.search(
        r"MULTI|SPLIT|CANALIZZ|CASSETTA|BTU|R-?32|PARETE", text
    ):
        return "CLIMA", "EXPLICIT_AIR_AIR_UNIT"
    if "SISTEMA DI CONTROLLO REMOTO WI-FI PER LINEA COMMERCIALE MIDEA" in text:
        return "CLIMA", "EXPLICIT_CLIMATE_CONTROL"
    return "NON_CLIMA", "NO_CLIMATE_PRODUCT_EVIDENCE"


def freeze_input():
    workbook = load_workbook(EXCEL, read_only=True, data_only=False)
    records = []
    sheets = []
    for sheet in workbook.worksheets:
        headers = [cell.value for cell in sheet[1]]
        normalized = [str(value or "").strip().lower() for value in headers]
        title_idx = next((i for i, value in enumerate(normalized) if value in {"nome", "titolo", "title"}), None)
        mpn_idx = next((i for i, value in enumerate(normalized) if value == "mpn"), None)
        if title_idx is None:
            sheets.append({"sheet": sheet.title, "processed": False, "headers": headers})
            continue
        sheet_count = 0
        for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            title = str(values[title_idx] or "").strip()
            if not title:
                continue
            scope, reason = classify_scope(title)
            current_mpn = values[mpn_idx] if mpn_idx is not None and mpn_idx < len(values) else None
            original = {
                str(headers[i] or f"column_{i + 1}"): value
                for i, value in enumerate(values)
                if value is not None
            }
            records.append({
                "row_id": f"{sheet.title}!{excel_row}",
                "sheet": sheet.title,
                "excel_row": excel_row,
                "title": title,
                "current_mpn": current_mpn,
                "reference": None,
                "brand_original": None,
                "benchmark_scope": scope,
                "scope_classification_reason": reason,
                "original_columns": original,
            })
            sheet_count += 1
        sheets.append({
            "sheet": sheet.title,
            "processed": True,
            "headers": headers,
            "data_rows": sheet_count,
            "title_column": headers[title_idx],
            "mpn_column": headers[mpn_idx] if mpn_idx is not None else None,
        })
    with open(FROZEN, "w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json_dump_line(record) + "\n")
    meta = {
        "excel_file": EXCEL.name,
        "excel_path": str(EXCEL),
        "excel_sha256": sha256_file(EXCEL),
        "frozen_input": str(FROZEN),
        "frozen_input_sha256": sha256_file(FROZEN),
        "excel_rows_total": len(records),
        "clima_rows": sum(record["benchmark_scope"] == "CLIMA" for record in records),
        "non_clima_rows": sum(record["benchmark_scope"] != "CLIMA" for record in records),
        "sheets": sheets,
    }
    return records, meta


def compact_relation(item):
    evidence = item.get("relation_evidence") or {}
    if evidence:
        return evidence.get("relation_type")
    evidences = item.get("relation_evidences") or []
    return evidences[0].get("relation_type") if evidences else None


def compact_component(item):
    slot = item.get("slot_id")
    requested_slot = item.get("taglia_btu")
    if requested_slot is None and isinstance(slot, str):
        match = re.search(r"slot_ui_(\d+)", slot)
        requested_slot = int(match.group(1)) if match else None
    return {
        "role": item.get("role") or ("UE" if item.get("is_ue") else "UI" if item.get("is_ui") else None),
        "requested_slot": requested_slot,
        "pt": str(item.get("code") or "") or None,
        "model": item.get("mfg_code"),
        "family": item.get("family_key") or item.get("catalog_family"),
        "variant": item.get("variant_full") or item.get("candidate_color"),
        "tier": item.get("evidence_tier"),
        "score": item.get("score"),
        "relation": compact_relation(item),
    }


def compact_top_candidates(pools):
    result = {}
    for slot, candidates in (pools or {}).items():
        if not candidates:
            continue
        result[slot] = [
            {
                "pt": str(item.get("code") or "") or None,
                "model": item.get("mfg_code"),
                "family": item.get("family_key") or item.get("catalog_family"),
                "variant": item.get("variant_full") or item.get("candidate_color"),
                "tier": item.get("evidence_tier"),
                "score": item.get("score"),
                "relation": compact_relation(item),
                "family_match": item.get("family_match"),
            }
            for item in candidates[:5]
        ]
    return result


def existing_mpn_status(current_mpn, automatic_mpn):
    current = str(current_mpn or "").strip()
    if not current or current.upper() in {"NT", "NON TROVATO", "N/D", "ND"}:
        return "NO_EXISTING_MPN"
    normalize = lambda value: re.sub(r"\s+", "", str(value or "").upper())
    return "SAME_AS_EXISTING" if normalize(current) == normalize(automatic_mpn) else "DIFFERENT_FROM_EXISTING"


def percentile(values, percent):
    if not values:
        return 0.0
    ordered = sorted(values)
    rank = (len(ordered) - 1) * percent
    lower = math.floor(rank)
    upper = math.ceil(rank)
    if lower == upper:
        return ordered[lower]
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (rank - lower)


def run_benchmark(records, meta):
    sys.path.insert(0, str(ROOT))
    os.chdir(ROOT)
    from catalog_search_engine import CatalogSearchEngine

    started = datetime.now(timezone.utc)
    engine = CatalogSearchEngine()
    engine._ensure_initialized()
    clima_records = [record for record in records if record["benchmark_scope"] == "CLIMA"]
    completed_ids = set()
    if OUTPUT.exists():
        with open(OUTPUT, "r", encoding="utf-8") as handle:
            for line in handle:
                try:
                    completed_ids.add(json.loads(line)["row_id"])
                except Exception:
                    continue
    pending = [record for record in clima_records if record["row_id"] not in completed_ids]
    mode = "a" if completed_ids else "w"
    with open(OUTPUT, mode, encoding="utf-8", newline="\n", buffering=1) as out:
        for index, record in enumerate(pending, start=1):
            wall_start = time.perf_counter()
            try:
                result = engine.search(record["title"], limit=20)
                qa = result.get("query_analysis") or {}
                context = qa.get("query_context") or {}
                bom = result.get("bom") or []
                selected_components = [compact_component(item) for item in bom]
                ue = next((item for item in selected_components if item.get("role") == "UE"), None)
                models = [item.get("model") for item in selected_components if item.get("model")]
                automatic_mpn = " + ".join(models) if models else None
                identity = result.get("product_identity_status") or qa.get("product_identity_status")
                if not selected_components:
                    main_status = "NO_MAIN_UNIT_SELECTED"
                elif identity:
                    main_status = identity
                elif qa.get("exact_token_candidates"):
                    main_status = "EXACT"
                else:
                    main_status = "DISCOVERY_ONLY"
                elapsed_ms = round((time.perf_counter() - wall_start) * 1000, 2)
                output_record = {
                    **{key: record.get(key) for key in (
                        "row_id", "sheet", "excel_row", "title", "current_mpn", "reference", "brand_original"
                    )},
                    "benchmark_scope": "CLIMA",
                    "detected_brand": result.get("detected_brand"),
                    "product_scope": result.get("product_scope") or context.get("product_scope") or "GENERAL",
                    "requested_btu": context.get("requested_btus") or [],
                    "requested_type": context.get("explicit_tipologia") or context.get("probable_tipologia"),
                    "requested_family": context.get("requested_family_key") or context.get("requested_family"),
                    "requested_color": context.get("query_color") or context.get("requested_color"),
                    "exact_tokens": qa.get("exact_token_candidates") or [],
                    "near_model_candidates": qa.get("near_model_candidates") or [],
                    "selected_components": selected_components,
                    "selected_ue": ({
                        "pt": ue.get("pt"), "model": ue.get("model"), "family": ue.get("family")
                    } if ue else None),
                    "automatic_mpn": automatic_mpn,
                    "mpn_final_allowed": bool(result.get("mpn_final_allowed")),
                    "main_unit_match_status": main_status,
                    "configuration_status": result.get("configuration_status") or "NOT_APPLICABLE",
                    "existing_mpn_comparison": existing_mpn_status(record.get("current_mpn"), automatic_mpn),
                    "top_candidates_by_slot": compact_top_candidates(result.get("candidate_pools")),
                    "runtime_ms": elapsed_ms,
                    "error": None,
                }
            except Exception as exc:
                elapsed_ms = round((time.perf_counter() - wall_start) * 1000, 2)
                output_record = {
                    **{key: record.get(key) for key in (
                        "row_id", "sheet", "excel_row", "title", "current_mpn", "reference", "brand_original"
                    )},
                    "benchmark_scope": "CLIMA",
                    "selected_components": [],
                    "selected_ue": None,
                    "automatic_mpn": None,
                    "main_unit_match_status": "RUNTIME_ERROR",
                    "configuration_status": "ERROR",
                    "existing_mpn_comparison": existing_mpn_status(record.get("current_mpn"), None),
                    "top_candidates_by_slot": {},
                    "runtime_ms": elapsed_ms,
                    "error": {"type": type(exc).__name__, "message": str(exc)},
                }
            out.write(json_dump_line(output_record) + "\n")
            total_done = len(completed_ids) + index
            if total_done % 25 == 0 or total_done == len(clima_records):
                progress = {
                    "completed": total_done,
                    "total": len(clima_records),
                    "last_row_id": record["row_id"],
                    "updated_at": datetime.now(timezone.utc).isoformat(),
                }
                PROGRESS.write_text(json.dumps(progress, indent=2), encoding="utf-8")
                print(json.dumps(progress), flush=True)

    results = []
    with open(OUTPUT, "r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                results.append(json.loads(line))
    runtimes = [record["runtime_ms"] for record in results]
    scope_counts = collections.Counter(record.get("product_scope") or "GENERAL" for record in results)
    normalized_scope_counts = collections.Counter()
    for key, value in scope_counts.items():
        normalized = key if key in {"MONOSPLIT", "MULTISPLIT", "UI_ONLY", "UE_ONLY"} else "GENERAL"
        normalized_scope_counts[normalized] += value
    for key in ("MONOSPLIT", "MULTISPLIT", "UI_ONLY", "UE_ONLY", "GENERAL"):
        normalized_scope_counts.setdefault(key, 0)
    identity_counts = collections.Counter(record.get("main_unit_match_status") for record in results)
    config_counts = collections.Counter(record.get("configuration_status") for record in results)
    comparison_counts = collections.Counter(record.get("existing_mpn_comparison") for record in results)
    summary = {
        **meta,
        "rows_processed": len(results),
        "rows_failed": sum(bool(record.get("error")) for record in results),
        "scope_counts": dict(normalized_scope_counts),
        "automatic_mpn_present": sum(bool(record.get("automatic_mpn")) for record in results),
        "automatic_mpn_missing": sum(not record.get("automatic_mpn") for record in results),
        "exact_identity_rows": identity_counts.get("EXACT", 0),
        "discovery_only_rows": identity_counts.get("DISCOVERY_ONLY", 0),
        "main_unit_status_counts": dict(identity_counts),
        "configuration_status_counts": dict(config_counts),
        "existing_mpn_comparison_counts": dict(comparison_counts),
        "runtime_total_seconds": round(sum(runtimes) / 1000, 3),
        "runtime_wall_seconds": round((datetime.now(timezone.utc) - started).total_seconds(), 3),
        "runtime_mean_ms": round(statistics.mean(runtimes), 3) if runtimes else 0,
        "runtime_median_ms": round(statistics.median(runtimes), 3) if runtimes else 0,
        "runtime_p95_ms": round(percentile(runtimes, 0.95), 3),
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary, started


def git_output(*args):
    return subprocess.check_output(["git", *args], cwd=ROOT, text=True).strip()


def build_manifest(meta, summary, started):
    runtime_paths = [
        "catalog_search_engine.py",
        "src/adapters/climate.py",
        "src/core/token_parser.py",
        "src/core/model_identity.py",
        "src/core/color_variants.py",
        "Knowledge/climatizzatori_compatibilita_master.json",
        "Knowledge/catalog_table_context.json",
    ]
    runtime_files = []
    for relative in runtime_paths:
        path = ROOT / relative
        if path.exists():
            runtime_files.append({"path": relative, "sha256": sha256_file(path)})
    manifest = {
        "phase": "BEFORE_V3",
        "git_commit": git_output("rev-parse", "HEAD"),
        "git_dirty": bool(git_output("status", "--short")),
        "excel_path": str(EXCEL),
        "excel_sha256": meta["excel_sha256"],
        "frozen_input": str(FROZEN),
        "frozen_input_sha256": meta["frozen_input_sha256"],
        "rows": summary["rows_processed"],
        "excel_rows_total": meta["excel_rows_total"],
        "clima_rows": meta["clima_rows"],
        "non_clima_rows": meta["non_clima_rows"],
        "runtime_files": runtime_files,
        "output_jsonl": str(OUTPUT),
        "output_jsonl_sha256": sha256_file(OUTPUT),
        "started_at": started.isoformat(),
        "completed_at": datetime.now(timezone.utc).isoformat(),
        "master_v3_active": False,
        "notes": "Baseline produced by the workspace matcher before clima_monosplit_light/clima_multisplit_light schema 3.0 integration.",
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--freeze-only", action="store_true")
    args = parser.parse_args()
    records, meta = freeze_input()
    print(json.dumps(meta, ensure_ascii=False, indent=2), flush=True)
    if args.freeze_only:
        return
    summary, started = run_benchmark(records, meta)
    manifest = build_manifest(meta, summary, started)
    print(json.dumps({"summary": summary, "manifest": manifest}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
