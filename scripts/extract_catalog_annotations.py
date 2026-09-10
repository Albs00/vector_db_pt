"""
Extract commercial annotations from the Puglia Termica PDF catalog layout.

Output datasets:
- Knowledge/catalog_annotations.json
- catalog_annotations_report.json
- catalog_annotations_review.xlsx
"""

from __future__ import annotations

import gzip
import hashlib
import json
import re
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Configure stdout
sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
LAYOUT_PATH = ROOT / "Knowledge" / "catalogo_layout.jsonl.gz"
TABLE_CONTEXT_PATH = ROOT / "Knowledge" / "catalog_table_context.json"
ROOT_TABLE_CONTEXT_PATH = ROOT / "catalog_table_context.json"

OUTPUT_JSON_PATH = ROOT / "Knowledge" / "catalog_annotations.json"
OUTPUT_REPORT_PATH = ROOT / "catalog_annotations_report.json"
OUTPUT_XLSX_PATH = ROOT / "catalog_annotations_review.xlsx"

def calc_table_id(page: int, title: str, bbox: list) -> str:
    signature = json.dumps(
        [int(page or 0), str(title or ""), [round(float(v), 2) for v in (bbox or [])]],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    digest = hashlib.sha1(signature.encode("utf-8")).hexdigest()[:10].upper()
    return f"PDF_P{int(page or 0):04d}_{digest}"

def is_red(hex_code: str) -> bool:
    if not hex_code or not hex_code.startswith("#") or len(hex_code) != 7:
        return False
    try:
        r = int(hex_code[1:3], 16)
        g = int(hex_code[3:5], 16)
        b = int(hex_code[5:7], 16)
        return r >= 150 and r > g * 1.35 and r > b * 1.35
    except Exception:
        return False

HEADER_EXCLUDES = {
    "CODICE", "MODELLO", "MODEL", "PREZZO", "PREZZO €", "PREZZO TOT €", "PREZZO UNITÀ €",
    "DISP", "KW", "CLASSE", "TUBI", "PESO", "TAGLIA", "MISURA", "GAS", "DATI ELETTRICI",
    "DATI IDRAULICI", "DESCRIZIONE", "DIAMETRO", "LUNGHEZZA", "LUNGHEZZA MT", "KW [F - C]",
    "KW [F-C]", "KW - CLASSE", "TUBI", "MODELLO KIT", "QUANTITÀ", "Q.TÀ", "INDICE FOTOGRAFICO",
    "INDICE DETTAGLIATO", "INDICE PER MARCHIO", "LISTINO 2026", "CATALOGO LISTINO 2026"
}

def is_filtered_out(text: str, y0: float, y1: float) -> bool:
    t = text.strip()
    up = t.upper()
    if not t or len(t) < 2:
        return True
    if y0 < 35 or y1 > 795:
        return True
    if t in ('✗', '•', '·', '*', '**', '***', 'F.E.S.'):
        return True
    if up in HEADER_EXCLUDES:
        return True
    if up.startswith("• CALDAIE") or up.startswith("• SISTEMI") or up.startswith("• SCALDABAGNI") or up.startswith("• VENTIL"):
        return True
    if re.fullmatch(r'\d+(?:[,\.]\d+)?\s*(?:mt|cm|mm|m|km|pz|kg|lt|l|bar|°C|/)?', t, re.IGNORECASE):
        return True
    if re.fullmatch(r'\d+/\d+', t):
        return True
    if re.fullmatch(r'Kw\s*\[.*?\](?:\s*-\s*[A-Z\+a-z]+)?', t, re.IGNORECASE):
        return True
    if re.fullmatch(r'[CF]?\s*\d+(?:[,\.]\d+)?(?:\s*-\s*[A-Z\+\d\.,/]+)?', t, re.IGNORECASE):
        return True
    if re.fullmatch(r'(?:A\+{0,3}|N\.D\.|n\.d\.)(?:\s+(?:A\+{0,3}|N\.D\.|n\.d\.))?', t):
        return True
    if re.fullmatch(r'N\.D\.\s+N\.D\.', t, re.IGNORECASE):
        return True
    return False

def split_clauses(text: str) -> list[str]:
    # If text is e.g. "telecomando incluso - griglia esclusa" or "Griglia esclusa, telecomando incluso"
    # we split them to keep clauses distinct.
    # Preserve full phrases like "Griglia esclusa da abbinare obbligatoriamente"
    if " - " in text:
        parts = [p.strip() for p in text.split(" - ") if p.strip()]
        if len(parts) > 1 and any(any(k in p.lower() for k in ["inclus", "esclus", "optional", "obblig"]) for p in parts):
            return parts
    if ", " in text:
        parts = [p.strip() for p in text.split(", ") if p.strip()]
        if len(parts) > 1 and any(any(k in p.lower() for k in ["inclus", "esclus", "optional", "obblig"]) for p in parts):
            return parts
    return [text]

def classify_annotation(raw_text: str) -> tuple[str, list[dict[str, str]]]:
    t = raw_text.strip()
    low = t.lower()
    
    # 1. COMPONENT_REQUIRED
    if any(w in low for w in ["obbligatoriamente", "obbligatorio da abbinare", "abbinare obbligatoriamente", "da acquistare obbligatoriamente", "acquisto obbligatorio"]):
        entities = []
        if "griglia" in low: entities.append({"type": "COMPONENT", "value": "GRIGLIA"})
        elif "macchina" in low: entities.append({"type": "COMPONENT", "value": "MACCHINA"})
        elif "comando" in low: entities.append({"type": "COMPONENT", "value": "COMANDO"})
        elif "bruciatore" in low: entities.append({"type": "COMPONENT", "value": "BRUCIATORE"})
        elif "fissaggio" in low: entities.append({"type": "COMPONENT", "value": "FISSAGGIO"})
        else: entities.append({"type": "OTHER", "value": "ABBINAMENTO OBBLIGATORIO"})
        return "COMPONENT_REQUIRED", entities

    # 2. PACKAGE_QUANTITY
    if re.search(r'(?:confez|scatol|pacc|blister).*?\bda\s+\d+', low) or re.search(r'vendibil.*?\bda\s+\d+', low):
        m = re.search(r'\bda\s+(\d+\s*(?:pezzi|pz|mt|metri|rotoli)?)\b', low)
        val = m.group(0).upper() if m else "CONFEZIONE"
        return "PACKAGE_QUANTITY", [{"type": "QUANTITY", "value": val}]

    # 3. SALE_QUANTITY
    if any(w in low for w in ["vendita solo a", "vendibile solo a", "vendita a rotolo", "vendita a barra", "solo a confezione", "solo a scatola", "solo a coppia", "solo a pezzo"]):
        val = "QUANTITA"
        if "rotolo" in low: val = "ROTOLO"
        elif "barra" in low: val = "BARRA"
        elif "confezion" in low: val = "CONFEZIONE"
        elif "scatola" in low: val = "SCATOLA"
        elif "coppia" in low: val = "COPPIA"
        elif "pezzo" in low: val = "PEZZO"
        return "SALE_QUANTITY", [{"type": "QUANTITY", "value": val}]

    # 4. COMPONENT_OPTIONAL
    if "optional" in low or "opzionale" in low:
        entities = []
        if "wifi" in low or "wi-fi" in low: entities.append({"type": "COMPONENT", "value": "WIFI"})
        elif "comando" in low: entities.append({"type": "COMPONENT", "value": "COMANDO"})
        elif "griglia" in low: entities.append({"type": "COMPONENT", "value": "GRIGLIA"})
        else:
            comp_val = re.sub(r'\(?\b(?:optional|opzionale)\b\)?', '', t, flags=re.IGNORECASE).strip() or "ACCESSORIO"
            entities.append({"type": "COMPONENT", "value": comp_val.upper()})
        return "COMPONENT_OPTIONAL", entities

    # 5. COMPONENT_EXCLUDED
    if any(w in low for w in ["escluso", "esclusa", "esclusi", "escluse", "non incluso", "non inclusa", "non inclusi", "non incluse", "privo di", "senza comando"]):
        entities = []
        if "griglia" in low: entities.append({"type": "COMPONENT", "value": "GRIGLIA"})
        if "telecomando" in low: entities.append({"type": "COMPONENT", "value": "TELECOMANDO"})
        elif "filo comando" in low or "comando a filo" in low or "filocomando" in low: entities.append({"type": "COMPONENT", "value": "COMANDO A FILO"})
        elif "comando" in low: entities.append({"type": "COMPONENT", "value": "COMANDO"})
        if "batteria" in low: entities.append({"type": "COMPONENT", "value": "BATTERIA"})
        if "struttura" in low or "fissaggio" in low: entities.append({"type": "COMPONENT", "value": "STRUTTURA DI FISSAGGIO"})
        if "piletta" in low: entities.append({"type": "COMPONENT", "value": "PILETTA"})
        if "sedile" in low: entities.append({"type": "COMPONENT", "value": "SEDILE"})
        if "centralina" in low: entities.append({"type": "COMPONENT", "value": "CENTRALINA"})
        if not entities: entities.append({"type": "COMPONENT", "value": "ACCESSORIO"})
        return "COMPONENT_EXCLUDED", entities

    # 6. FEATURE_INCLUDED
    if any(w in low for w in ["dotata di", "dotato di", "di serie", "integrato", "integrata", "con pompa", "con wifi", "con wi-fi", "ionizzatore"]):
        entities = []
        if "pompa" in low: entities.append({"type": "FEATURE", "value": "POMPA SCARICO CONDENSA"})
        if "ionizzatore" in low: entities.append({"type": "FEATURE", "value": "IONIZZATORE"})
        if "wifi" in low or "wi-fi" in low: entities.append({"type": "FEATURE", "value": "WIFI INTEGRATO"})
        if "antigelo" in low: entities.append({"type": "FEATURE", "value": "ANTIGELO"})
        if not entities: entities.append({"type": "FEATURE", "value": t.upper()})
        return "FEATURE_INCLUDED", entities

    # 7. COMPONENT_INCLUDED
    if any(w in low for w in ["incluso", "inclusa", "inclusi", "incluse", "compreso", "compresa", "compresi", "completa di", "completo di"]):
        entities = []
        if "telecomando" in low: entities.append({"type": "COMPONENT", "value": "TELECOMANDO"})
        if "comando a filo" in low or "filocomando" in low or "filo comando" in low: entities.append({"type": "COMPONENT", "value": "COMANDO A FILO"})
        elif "comando" in low and "telecomando" not in low: entities.append({"type": "COMPONENT", "value": "COMANDO"})
        if "pompa" in low and "condensa" in low: entities.append({"type": "COMPONENT", "value": "POMPA SCARICO CONDENSA"})
        if "griglia" in low: entities.append({"type": "COMPONENT", "value": "GRIGLIA"})
        if "piletta" in low: entities.append({"type": "COMPONENT", "value": "PILETTA"})
        if "valvola" in low: entities.append({"type": "COMPONENT", "value": "VALVOLA"})
        if not entities: entities.append({"type": "COMPONENT", "value": "ACCESSORIO"})
        return "COMPONENT_INCLUDED", entities

    # 8. INSTALLATION_REQUIREMENT
    if any(w in low for w in ["avviamento", "prima accensione", "installazione", "fissaggio a", "montaggio"]):
        val = t.upper()
        if "avviamento" in low: val = "AVVIAMENTO"
        elif "accensione" in low: val = "PRIMA ACCENSIONE"
        elif "installazione" in low: val = "INSTALLAZIONE"
        return "INSTALLATION_REQUIREMENT", [{"type": "FEATURE", "value": val}]

    # 9. COMPATIBILITY_NOTE
    if any(w in low for w in ["compatib", "abbinabile", "abbinare", "abbinamento", "verificare la compatibilità"]):
        return "COMPATIBILITY_NOTE", [{"type": "OTHER", "value": "COMPATIBILITA"}]

    # 10. USAGE_LIMITATION
    if any(w in low for w in ["solo per", "solo uso", "solo riscaldamento", "può essere montato solo", "vedere ultima pagina"]):
        return "USAGE_LIMITATION", [{"type": "OTHER", "value": t.upper()}]

    # 11. FEATURE
    if any(w in low for w in ["reversibile", "trifase", "monofase", "alimentare", "produzione", "wi-fi", "wifi"]):
        val = t.upper()
        if "reversibile" in low: val = "REVERSIBILE" if "non" not in low else "NON REVERSIBILE"
        elif "trifase" in low: val = "TRIFASE"
        return "FEATURE", [{"type": "FEATURE", "value": val}]

    return "OTHER", [{"type": "OTHER", "value": t.upper()}]

commercial_indicators = [
    "incluso", "inclusa", "inclusi", "incluse",
    "compreso", "compresa", "compresi", "comprese",
    "completo di", "completa di",
    "escluso", "esclusa", "esclusi", "escluse",
    "senza ", "privo di", "non inclus",
    "obbligatori", "da abbinare", "abbinare", "da completare", "da acquistare",
    "optional", "opzional",
    "di serie", "integrat", "dotat",
    "vendita solo", "vendibile", "confezion", "scatola da", "solo a rotolo", "solo a barra", "solo a scatola",
    "prima accensione", "avviamento",
    "a noleggio", "anni di garanzia",
    "compatibil"
]

def is_commercial_note(text: str, has_red: bool) -> bool:
    low = text.lower()
    if has_red:
        if any(k in low for k in commercial_indicators):
            return True
        if any(w in low for w in ["reversibile", "trifase", "monofase", "alimentare", "promozioni", "preventivi", "normativa"]):
            return True
        return False
    else:
        return any(k in low for k in commercial_indicators)

def load_tables() -> dict[int, dict[str, dict[str, Any]]]:
    tables_by_page = defaultdict(dict)
    
    # 1. From root catalog_table_context.json (has layout bounding boxes)
    if ROOT_TABLE_CONTEXT_PATH.exists():
        with ROOT_TABLE_CONTEXT_PATH.open("r", encoding="utf-8") as f:
            root_data = json.load(f)
        for code, rec in root_data.items():
            page = int(rec.get("page") or 0)
            title = rec.get("table_title") or ""
            bbox = (rec.get("layout_evidence") or {}).get("table_bbox") or []
            title_bbox = (rec.get("layout_evidence") or {}).get("title_bbox") or []
            code_bbox = (rec.get("layout_evidence") or {}).get("code_bbox") or []
            section = rec.get("section_title") or ""
            tid = calc_table_id(page, title, bbox)
            if tid not in tables_by_page[page]:
                tables_by_page[page][tid] = {
                    "table_id": tid,
                    "page": page,
                    "title": title,
                    "section": section,
                    "table_bbox": bbox,
                    "title_bbox": title_bbox,
                    "min_code_y": code_bbox[1] if code_bbox else (bbox[3] if bbox else 0),
                    "max_code_y": code_bbox[3] if code_bbox else (bbox[3] if bbox else 0),
                    "min_x": bbox[0] if bbox else 0,
                    "max_x": bbox[2] if bbox else 600,
                }
            if code_bbox:
                t = tables_by_page[page][tid]
                t["min_code_y"] = min(t["min_code_y"], code_bbox[1])
                t["max_code_y"] = max(t["max_code_y"], code_bbox[3])
                t["min_x"] = min(t["min_x"], code_bbox[0])
                t["max_x"] = max(t["max_x"], code_bbox[2])

            # Alternates in root
            for alt in rec.get("alternate_table_contexts") or []:
                apage = int(alt.get("page") or 0)
                atitle = alt.get("table_title") or ""
                abbox = alt.get("table_bbox") or []
                atid = calc_table_id(apage, atitle, abbox)
                if apage and atid not in tables_by_page[apage]:
                    tables_by_page[apage][atid] = {
                        "table_id": atid,
                        "page": apage,
                        "title": atitle,
                        "section": section,
                        "table_bbox": abbox,
                        "title_bbox": [],
                        "min_code_y": abbox[1] if abbox else 0,
                        "max_code_y": abbox[3] if abbox else 800,
                        "min_x": abbox[0] if abbox else 0,
                        "max_x": abbox[2] if abbox else 600,
                    }

    # 2. Complement from Knowledge/catalog_table_context.json
    if TABLE_CONTEXT_PATH.exists():
        with TABLE_CONTEXT_PATH.open("r", encoding="utf-8") as f:
            k_data = json.load(f)
        for code, rec in k_data.items():
            tid = rec.get("table_id")
            page = int(rec.get("page") or 0)
            if tid and page:
                if tid not in tables_by_page[page]:
                    tables_by_page[page][tid] = {
                        "table_id": tid,
                        "page": page,
                        "title": rec.get("table_title") or "",
                        "section": rec.get("section_title") or "",
                        "table_bbox": [],
                        "title_bbox": [],
                        "min_code_y": 0,
                        "max_code_y": 800,
                        "min_x": 0,
                        "max_x": 600,
                    }
            for alt in rec.get("alternate_table_contexts") or []:
                apage = int(alt.get("page") or 0)
                atid = alt.get("table_id")
                atitle = alt.get("table_title") or ""
                if apage and atid:
                    if atid not in tables_by_page[apage]:
                        tables_by_page[apage][atid] = {
                            "table_id": atid,
                            "page": apage,
                            "title": atitle,
                            "section": "",
                            "table_bbox": [],
                            "title_bbox": [],
                            "min_code_y": 0,
                            "max_code_y": 800,
                            "min_x": 0,
                            "max_x": 600,
                        }

    return tables_by_page

def match_table_for_bbox(
    page_num: int, bbox: list[float], tables_by_page: dict[int, dict[str, dict[str, Any]]], default_section: str = ""
) -> tuple[Optional[dict[str, Any]], str, str, float]:
    tables = list(tables_by_page.get(page_num, {}).values())
    if not tables:
        return None, "TABLE", default_section, 0.70
    if len(tables) == 1:
        t = tables[0]
        return t, "TABLE", t.get("section") or default_section, 0.95
    
    x0, y0, x1, y1 = bbox
    cx = (x0 + x1) / 2
    cy = (y0 + y1) / 2
    
    candidates = []
    for t in tables:
        t_box = t["table_bbox"]
        title_box = t["title_bbox"]
        
        col_x0 = min(t["min_x"], t_box[0] if t_box else 0, title_box[0] if title_box else 0) - 15
        col_x1 = max(t["max_x"], t_box[2] if t_box else 600, title_box[2] if title_box else 600) + 15
        
        if col_x0 <= cx <= col_x1:
            top_y = (title_box[1] if title_box else (t_box[1] - 20 if t_box else 0)) - 5
            header_y = (t_box[3] if t_box else t["min_code_y"]) + 10
            bottom_y = t["max_code_y"] + 30
            
            if top_y <= cy <= header_y:
                dist = abs(cy - (title_box[3] if title_box else (t_box[1] if t_box else cy)))
                candidates.append((dist, 1.0, "TABLE", t))
            elif header_y <= cy <= bottom_y:
                dist = abs(cy - t["min_code_y"])
                candidates.append((dist + 100, 0.95, "ROW", t))
            elif cy > bottom_y:
                dist = cy - bottom_y
                candidates.append((dist + 500, 0.85, "TABLE", t))
            else:
                dist = top_y - cy
                candidates.append((dist + 1000, 0.70, "TABLE", t))
                
    if candidates:
        candidates.sort(key=lambda x: x[0])
        best = candidates[0]
        return best[3], best[2], best[3].get("section") or default_section, best[1]
        
    best_t = None
    min_d = float('inf')
    for t in tables:
        t_cx = (t["min_x"] + t["max_x"]) / 2
        t_cy = (t["min_code_y"] + (t["table_bbox"][1] if t["table_bbox"] else 0)) / 2
        d = (cx - t_cx)**2 + (cy - t_cy)**2
        if d < min_d:
            min_d = d
            best_t = t
    return best_t, "TABLE", best_t.get("section") if best_t else default_section, 0.75

def run_extraction() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    print("Starting commercial annotation extraction from catalog layout...")
    t0 = time.time()
    tables_by_page = load_tables()
    print(f"Loaded table contexts for {len(tables_by_page)} pages.")

    annotations: list[dict[str, Any]] = []
    seen_annotations: set[tuple[int, str, float, float]] = set()

    total_pages_scanned = 0
    with gzip.open(LAYOUT_PATH, "rt", encoding="utf-8") as f:
        for line in f:
            total_pages_scanned += 1
            data = json.loads(line)
            page_num = int(data.get("page") or 0)

            # Discover section title from prominent text on page
            page_section = ""
            for b in data.get("blocks", []):
                for l in b.get("lines", []):
                    for s in l.get("spans", []):
                        if s.get("size_pt", 0) >= 12.0 and s.get("bbox", [0, 0, 0, 0])[1] < 40:
                            page_section = s.get("text", "").strip()
                            break
                    if page_section:
                        break
                if page_section:
                    break

            # Skip non-product introductory pages (legal terms, general index)
            if page_num < 28 or "INDICE" in page_section.upper() or "CONDIZIONI" in page_section.upper():
                continue

            blocks = data.get("blocks", [])
            # Collect lines on this page
            page_lines = []
            for b in blocks:
                for l in b.get("lines", []):
                    spans = l.get("spans", [])
                    txt = "".join(s.get("text", "") for s in spans).strip()
                    bbox = l.get("bbox", [0, 0, 0, 0])
                    page_lines.append((txt, bbox, spans))

            # 1. First pass: check for row-level combinations (e.g. Kit WI-FI + OPTIONAL)
            # Find lines that are "OPTIONAL" or "optional" and look for adjacent accessory description on same y
            optional_lines = [item for item in page_lines if item[0].lower() in ("optional", "(optional)", "opzionale")]
            for opt_txt, opt_box, opt_spans in optional_lines:
                opt_y = (opt_box[1] + opt_box[3]) / 2
                # Look for companion text on the same row (+- 6pt)
                for comp_txt, comp_box, comp_spans in page_lines:
                    if comp_txt == opt_txt:
                        continue
                    comp_y = (comp_box[1] + comp_box[3]) / 2
                    if abs(opt_y - comp_y) <= 6.0 and abs(comp_box[0] - opt_box[0]) > 20.0:
                        low_c = comp_txt.lower()
                        if any(w in low_c for w in ["wifi", "wi-fi", "griglia", "filo comando", "comando a filo", "comando"]):
                            combined_text = f"{comp_txt} optional"
                            comb_bbox = [
                                min(opt_box[0], comp_box[0]),
                                min(opt_box[1], comp_box[1]),
                                max(opt_box[2], comp_box[2]),
                                max(opt_box[3], comp_box[3]),
                            ]
                            dedup_key = (page_num, combined_text.upper(), round(comb_bbox[0], 1), round(comb_bbox[1], 1))
                            if dedup_key not in seen_annotations:
                                seen_annotations.add(dedup_key)
                                matched_table, scope, sec, conf = match_table_for_bbox(
                                    page_num, comb_bbox, tables_by_page, default_section=page_section
                                )
                                atype, entities = classify_annotation(combined_text)
                                ann = {
                                    "annotation_id": f"ANN_P{page_num:04d}_{len(annotations)+1:04d}",
                                    "page": page_num,
                                    "table_id": matched_table["table_id"] if matched_table else None,
                                    "table_title": matched_table["title"] if matched_table else None,
                                    "section_title": sec,
                                    "scope": "ACCESSORY_BLOCK" if (matched_table and "ACCESSORI" in (matched_table.get("title") or "").upper()) else "ROW",
                                    "raw_text": combined_text,
                                    "bbox": {
                                        "x0": round(comb_bbox[0], 2),
                                        "y0": round(comb_bbox[1], 2),
                                        "x1": round(comb_bbox[2], 2),
                                        "y1": round(comb_bbox[3], 2),
                                    },
                                    "source": "PDF_LAYOUT",
                                    "confidence": conf,
                                    "annotation_type": atype,
                                    "entities": entities,
                                }
                                annotations.append(ann)

            # 2. Main pass: scan lines for red annotations or commercial notes
            for line_text, bbox, spans in page_lines:
                if is_filtered_out(line_text, bbox[1], bbox[3]):
                    continue

                has_red = any(is_red(s.get("color")) for s in spans)
                has_comm = is_commercial_note(line_text, has_red)

                # Check if it's an editorial subtitle under a table title
                is_subtitle = False
                for t in tables_by_page.get(page_num, {}).values():
                    if t["title_bbox"]:
                        if abs(bbox[0] - t["title_bbox"][0]) < 25 and t["title_bbox"][3] <= bbox[1] <= t["title_bbox"][3] + 25:
                            # Subtitle immediately under title
                            if len(line_text.split()) >= 2 and not is_filtered_out(line_text, bbox[1], bbox[3]):
                                is_subtitle = True
                                break

                if not (has_comm or (has_red and is_subtitle)):
                    continue

                clauses = split_clauses(line_text)
                for clause in clauses:
                    if is_filtered_out(clause, bbox[1], bbox[3]):
                        continue
                    dedup_key = (page_num, clause.upper(), round(bbox[0], 1), round(bbox[1], 1))
                    if dedup_key in seen_annotations:
                        continue
                    seen_annotations.add(dedup_key)

                    matched_table, scope, sec, conf = match_table_for_bbox(
                        page_num, bbox, tables_by_page, default_section=page_section
                    )
                    atype, entities = classify_annotation(clause)

                    matched_table_title = (matched_table.get("title") or "") if matched_table else ""
                    # Adjust scope if clearly a table-level note or accessory note
                    if is_subtitle:
                        scope = "TABLE"
                        conf = max(conf, 1.0)
                    elif "ACCESSORI" in matched_table_title.upper():
                        scope = "ACCESSORY_BLOCK"
                    elif atype in ("SALE_QUANTITY", "PACKAGE_QUANTITY"):
                        scope = "ROW"

                    ann = {
                        "annotation_id": f"ANN_P{page_num:04d}_{len(annotations)+1:04d}",
                        "page": page_num,
                        "table_id": matched_table["table_id"] if matched_table else None,
                        "table_title": matched_table["title"] if matched_table else None,
                        "section_title": sec,
                        "scope": scope,
                        "raw_text": clause,
                        "bbox": {
                            "x0": round(bbox[0], 2),
                            "y0": round(bbox[1], 2),
                            "x1": round(bbox[2], 2),
                            "y1": round(bbox[3], 2),
                        },
                        "source": "PDF_LAYOUT",
                        "confidence": conf,
                        "annotation_type": atype,
                        "entities": entities,
                    }
                    annotations.append(ann)

    elapsed = time.time() - t0
    print(f"Extraction completed in {elapsed:.2f}s across {total_pages_scanned} pages.")
    print(f"Total annotations extracted: {len(annotations)}")

    # Audit report preparation
    pages_involved = sorted(list(set(a["page"] for a in annotations)))
    tables_involved = sorted(list(set(a["table_id"] for a in annotations if a.get("table_id"))))
    type_distribution = Counter(a["annotation_type"] for a in annotations)
    text_distribution = Counter(a["raw_text"] for a in annotations)
    top_100_texts = [{"text": t, "count": c} for t, c in text_distribution.most_common(100)]
    without_table_id = [a["annotation_id"] for a in annotations if not a.get("table_id")]
    low_confidence = [a["annotation_id"] for a in annotations if a["confidence"] < 0.8]

    report = {
        "total_annotations": len(annotations),
        "pages_count": len(pages_involved),
        "tables_count": len(tables_involved),
        "annotation_type_distribution": dict(type_distribution.most_common()),
        "top_100_frequent_texts": top_100_texts,
        "annotations_without_table_id_count": len(without_table_id),
        "annotations_without_table_id_sample": without_table_id[:50],
        "annotations_low_confidence_count": len(low_confidence),
        "annotations_low_confidence_sample": low_confidence[:50],
    }

    return annotations, report

def save_outputs(annotations: list[dict[str, Any]], report: dict[str, Any]):
    print(f"Writing {OUTPUT_JSON_PATH}...")
    with OUTPUT_JSON_PATH.open("w", encoding="utf-8") as f:
        json.dump(annotations, f, indent=2, ensure_ascii=False)

    print(f"Writing {OUTPUT_REPORT_PATH}...")
    with OUTPUT_REPORT_PATH.open("w", encoding="utf-8") as f:
        json.dump(report, f, indent=2, ensure_ascii=False)

    print(f"Writing {OUTPUT_XLSX_PATH}...")
    # Manual validation Excel:
    # Columns: page, table_id, table_title, raw_text, annotation_type, entities, confidence, review_status
    # Sort order: first confidence ascending (lowest confidence first), then frequency descending
    text_freq = Counter(a["raw_text"] for a in annotations)
    sorted_annotations = sorted(
        annotations,
        key=lambda a: (a["confidence"], -text_freq[a["raw_text"]], a["page"])
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalog Annotations"

    # Header styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    headers = [
        "page", "table_id", "table_title", "raw_text", "annotation_type", "entities", "confidence", "review_status"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for a in sorted_annotations:
        entities_str = ", ".join(f"{e['type']}:{e['value']}" for e in a.get("entities", []))
        review_status = "VERIFIED" if a["confidence"] >= 1.0 else "PENDING_REVIEW"
        row = [
            a["page"],
            a.get("table_id") or "",
            a.get("table_title") or "",
            a["raw_text"],
            a["annotation_type"],
            entities_str,
            round(a["confidence"], 2),
            review_status,
        ]
        ws.append(row)

    # Apply widths and borders
    col_widths = {
        "A": 8,   # page
        "B": 24,  # table_id
        "C": 35,  # table_title
        "D": 45,  # raw_text
        "E": 28,  # annotation_type
        "F": 30,  # entities
        "G": 12,  # confidence
        "H": 18,  # review_status
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Freeze header
    ws.freeze_panes = "A2"

    wb.save(OUTPUT_XLSX_PATH)
    print(f"Successfully generated {OUTPUT_XLSX_PATH} with {len(sorted_annotations)} rows.")

if __name__ == "__main__":
    annotations, report = run_extraction()
    save_outputs(annotations, report)
    print("\n--- EXTRACTION REPORT SUMMARY ---")
    print(f"Total annotations: {report['total_annotations']}")
    print(f"Pages involved: {report['pages_count']}")
    print(f"Tables involved: {report['tables_count']}")
    print("Annotation type distribution:")
    for k, v in report['annotation_type_distribution'].items():
        print(f"  {k:26s}: {v:5d}")
    print(f"Annotations without table_id: {report['annotations_without_table_id_count']}")
    print(f"Annotations low confidence (< 0.8): {report['annotations_low_confidence_count']}")
