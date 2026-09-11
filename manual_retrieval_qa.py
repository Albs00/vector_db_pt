#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
manual_retrieval_qa.py
Strumento di Quality Assurance (QA) per test umano del motore di candidate retrieval.

Supporta:
1. Modalità Interattiva / Singola Query:
   Input:
     - Titolo query
     - Riferimento opzionale (es. codice fornitore, riga)
     - MPN attuale opzionale
   Output:
     - Query Analysis (Brand rilevato, Dominio/Adapter, Token esatti, Near-model, Contesto)
     - Top 5 candidati per ciascun candidate_pools[slot_id]
       (Codice PT, Nome catalogo, MFG code, Match Type, Identity Evidence,
        Relation Evidence separata, Relation Type, Pagina/Provenance, Rank within-slot)
     - BOM Candidata Completa assemblata

2. Modalità Batch (XLSX / CSV):
   Input: File Excel (.xlsx) o CSV con colonne (titolo/descrizione/query, riferimento, mpn)
   Output: File Excel (.xlsx) o CSV con l'intera analisi di retrieval e la BOM candidata.

Uso:
  python manual_retrieval_qa.py --title "Daikin Climatizzatore Monosplit Bluevolution Emura III 12000 btu FTXJ35AB9"
  python manual_retrieval_qa.py --batch dataset_holdout.xlsx --output results_qa.xlsx
  python manual_retrieval_qa.py --interactive
"""

import sys
import os
import re
import csv
import json
import argparse
from pathlib import Path
from typing import Dict, Any, List, Optional

# Aggiungi root del progetto al path
BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR))

from catalog_search_engine import CatalogSearchEngine


def format_qa_single_report(
    title: str,
    search_res: Dict[str, Any],
    reference: Optional[str] = None,
    current_mpn: Optional[str] = None
) -> str:
    lines = []
    sep = "=" * 80
    subsep = "-" * 80

    lines.append(sep)
    lines.append("  MANUAL RETRIEVAL QA REPORT")
    lines.append(sep)
    lines.append(f"TITOLO QUERY : {title}")
    if reference:
        lines.append(f"RIFERIMENTO  : {reference}")
    if current_mpn:
        lines.append(f"MPN ATTUALE  : {current_mpn}")
    lines.append(f"TEMPO ESEC.  : {search_res.get('execution_time_ms', 0)} ms")

    # 1. QUERY ANALYSIS
    lines.append(subsep)
    lines.append("1. QUERY ANALYSIS")
    lines.append(subsep)
    lines.append(f"  • Detected Brand     : {search_res.get('detected_brand') or 'N/D'}")
    lines.append(f"  • Category / Domain  : {search_res.get('detected_category') or 'CLIMA'}")
    lines.append(f"  • Match Type Global  : {search_res.get('match_type') or 'DESCRIPTIVE'}")

    qa_info = search_res.get("query_analysis", {})
    q_ctx = qa_info.get("query_context", {})
    if q_ctx:
        btus = q_ctx.get("requested_btus", [])
        lines.append(f"  • BTU Richiesti      : {btus if btus else 'Non specificati'}")
        p_scope = search_res.get("product_scope") or q_ctx.get("product_scope")
        lines.append(f"  • Product Scope      : {p_scope or 'N/D'}")
        lines.append(f"  • Configurazione     : {'Monosplit' if q_ctx.get('is_monosplit') else ('Multisplit' if q_ctx.get('is_multisplit') else ('UI Only' if q_ctx.get('is_ui_only') else ('UE Only' if q_ctx.get('is_ue_only') else 'Singola/Generale')))}")
        c_status = search_res.get("compatibility_status") or qa_info.get("compatibility_status")
        if c_status:
            lines.append(f"  • Compatibility Stat : {c_status}")
        lines.append(f"  • Tipologia Esplicita: {q_ctx.get('explicit_tipologia') or 'Nessuna (no hard filter)'}")
        lines.append(f"  • Tipologia Probabile: {q_ctx.get('probable_tipologia') or 'Non determinata'}")
        lines.append(f"  • Fase Richiesta     : {q_ctx.get('phase') or 'Non specificata'}")
        if q_ctx.get("feature_wifi"):
            lines.append(f"  • Feature Wi-Fi      : {q_ctx.get('feature_wifi')}")
        req_fam_key = qa_info.get("requested_family_key")
        lines.append(f"  • Family Key Req     : {req_fam_key if req_fam_key is not None else 'null'}")
        lines.append(f"  • Reliable Anchor    : {q_ctx.get('has_reliable_anchor', False)}")

    exact_cands = qa_info.get("exact_token_candidates", [])
    lines.append(f"  • Exact Tokens Trovati ({len(exact_cands)}):")
    if exact_cands:
        for ec in exact_cands:
            lines.append(f"      - PT: {ec.get('code')} | MFG: {ec.get('mfg_code')} | Tipo: {ec.get('match_type')}")
    else:
        lines.append("      - Nessun match esatto su codice produttore o anagrafica catalogo")

    near_cands = qa_info.get("near_model_candidates", [])
    if near_cands:
        lines.append(f"  • Near-Model Candidates ({len(near_cands)}):")
        for nc in near_cands[:5]:
            lines.append(f"      - PT: {nc.get('code')} | MFG: {nc.get('mfg_code')}")

    # 2. CANDIDATE POOLS (Top 5 per ciascun slot_id)
    lines.append(subsep)
    lines.append("2. CANDIDATE POOLS PER SLOT (Top 5)")
    lines.append(subsep)

    pools = search_res.get("candidate_pools", {})
    if not pools:
        lines.append("  Nessun candidate pool generato.")
    else:
        for slot_id, cands in pools.items():
            lines.append(f"\n[SLOT: {slot_id.upper()}] - Totale candidati: {len(cands)}")
            lines.append(f"{'#':<3} {'CODICE PT':<10} {'MFG CODE':<16} {'FAMILY KEY':<24} {'F.MATCH':<8} {'TIER':<6} {'SCORE':<7} {'RELATION TYPE':<24} {'PAG.':<5} {'NOME CATALOGO'}")
            lines.append("-" * 145)

            for rank_idx, c in enumerate(cands[:5], 1):
                pt_code = str(c.get("code") or "")
                mfg = str(c.get("mfg_code") or "-")[:15]
                fam_key = str(c.get("family_key") or "-")[:23]
                fam_match = str(c.get("family_match") or c.get("table_family_match") or "-")[:7]
                tier = f"T{c.get('evidence_tier', 9)}"
                score = f"{c.get('score', 0.0):.1f}"

                rel_meta = c.get("relation_evidence")
                rel_evs = c.get("relation_evidences") or []
                if rel_meta:
                    rel_str = f"{rel_meta.get('relation_type')} ({rel_meta.get('relation_strength')})"
                    if len(rel_evs) > 1:
                        rel_str += f" [+{len(rel_evs)-1}]"
                elif rel_evs:
                    rel_str = f"{rel_evs[0].get('relation_type')}"
                    if len(rel_evs) > 1:
                        rel_str += f" [+{len(rel_evs)-1}]"
                else:
                    rel_str = "-"
                rel_str = rel_str[:23]

                pag = str(c.get("primary_page") or "-")[:4]
                name = str(c.get("name") or "")[:40]

                lines.append(f"{rank_idx:<3} {pt_code:<10} {mfg:<16} {fam_key:<24} {fam_match:<8} {tier:<6} {score:<7} {rel_str:<24} {pag:<5} {name}")

    # 3. BOM CANDIDATA COMPLETA
    lines.append(subsep)
    lines.append("3. BOM CANDIDATA RECUPERATA")
    lines.append(subsep)
    lines.append(f"  • QUERY_COLOR          : {q_ctx.get('query_color') or '-'}")
    lines.append("")

    # Identifica i componenti primari della BOM dal motore o fallback
    bom_items = []
    engine_bom = search_res.get("bom")
    if engine_bom:
        for it in engine_bom:
            role_lbl = it.get("role_label") or it.get("role") or "Componente BOM"
            bom_items.append((role_lbl, it))
    else:
        # Fallback slot-based
        ue_cands = pools.get("slot_ue", [])
        if ue_cands:
            bom_items.append(("UE (Motore Esterno)", ue_cands[0]))
        if q_ctx and q_ctx.get("requested_btus"):
            for btu in q_ctx.get("requested_btus"):
                slot_name = f"slot_ui_{btu}"
                ui_cands = pools.get(slot_name, pools.get("slot_ui", []))
                if ui_cands:
                    bom_items.append((f"UI {btu} BTU", ui_cands[0]))
        else:
            ui_cands = pools.get("slot_ui", [])
            if ui_cands:
                bom_items.append(("UI (Unità Interna)", ui_cands[0]))

    if not bom_items:
        results = search_res.get("results", [])
        for idx, r in enumerate(results[:2], 1):
            bom_items.append((f"Candidato #{idx}", r))

    for role_lbl, item in bom_items:
        pt_code = item.get("code")
        name = item.get("name")
        mfg = item.get("mfg_code") or "-"
        score = item.get("score")
        tier = item.get("evidence_tier")

        ident_meta = item.get("identity_evidence") or {}
        rel_meta = item.get("relation_evidence") or {}
        rel_evs = item.get("relation_evidences") or []

        ident_desc = f"Tier {tier} [{ident_meta.get('match_type', 'N/D')}]"
        if rel_meta:
            rel_desc = f"{rel_meta.get('relation_type')} (conf: {rel_meta.get('confidence')}, prov: {rel_meta.get('provenance')})"
        else:
            rel_desc = "Nessuna relazione (riconoscimento lessicale/diretto)"

        lines.append(f"  • {role_lbl:<22}: PT {pt_code} | MFG: {mfg}")
        lines.append(f"    Nome Catalogo       : {name}")
        lines.append(f"    Slot Naturale       : {item.get('slot_id', '-')}")
        lines.append(f"    Family Context      : Key: {item.get('family_key', '-')}, Match: {item.get('family_match', '-')}, Family: {item.get('catalog_family', '-')}")
        lines.append(f"    Identity Evidence   : {ident_desc}")
        lines.append(f"    Relation Evidence   : {rel_desc}")
        lines.append(f"    CANDIDATE_COLOR     : {item.get('candidate_color') or '-'}")
        lines.append(f"    VARIANT_FULL        : {item.get('variant_full') or '-'}")
        lines.append(f"    VARIANT_CONFLICT    : {bool(item.get('variant_conflict'))}")
        if rel_evs and len(rel_evs) > 1:
            rel_summary = ", ".join(f"{r.get('relation_type')} (src:{r.get('source_code')})" for r in rel_evs)
            lines.append(f"    Relation Evidences  : {rel_summary}")
        lines.append(f"    Score Complessivo   : {score}")
        lines.append("")

    # 4. INFORMAZIONI DI COMPATIBILITÀ (Separate dalla BOM)
    lines.append(subsep)
    lines.append("4. INFORMAZIONI DI COMPATIBILITÀ (Separate dalla BOM)")
    lines.append(subsep)
    lines.append(f"  • Product Scope        : {search_res.get('product_scope') or 'N/D'}")
    lines.append(f"  • Compatibility Status : {search_res.get('compatibility_status') or 'NOT_APPLICABLE'}")
    lines.append(f"  • PAIRING_REASON        : {search_res.get('pairing_reason') or 'NOT_APPLICABLE'}")
    lines.append(f"  • PAIRING_STATUS        : {search_res.get('pairing_status') or 'NOT_APPLICABLE'}")
    lines.append(f"  • PRODUCT_IDENTITY_STATUS: {search_res.get('product_identity_status') or 'NOT_APPLICABLE'}")
    lines.append(f"  • CONFIGURATION_STATUS : {search_res.get('configuration_status') or 'NOT_APPLICABLE'}")
    lines.append(f"  • MPN Final Allowed     : {bool(search_res.get('mpn_final_allowed'))}")
    pairing_breakdown = search_res.get("pairing_score_breakdown") or {}
    if pairing_breakdown:
        ordered_pairing_keys = ["PDF_PAIRING", "EXPLICIT_MODEL", "EXPLICIT_MODEL_REVISION", "FAMILY", "MASTER", "MASTER_PAIRWISE_SELECTED_UI", "BTU", "VARIANT", "FINAL"]
        breakdown_text = ", ".join(
            f"{key}={pairing_breakdown.get(key)}"
            for key in ordered_pairing_keys
            if key in pairing_breakdown
        )
        lines.append(f"  • PAIRING_SCORE_BREAKDOWN: {breakdown_text}")
    pairing_diag = search_res.get("pairing_diagnostics") or {}
    lines.append(f"  • EXPLICIT_UE_TOKEN     : {pairing_diag.get('EXPLICIT_UE_TOKEN') or '-'}")
    lines.append(f"  • EXPLICIT_UE_MATCH_TYPE: {pairing_diag.get('EXPLICIT_UE_MATCH_TYPE') or '-'}")
    lines.append(
        "  • EXPLICIT_UE_CANDIDATES: "
        f"{json.dumps(pairing_diag.get('EXPLICIT_UE_CANDIDATES') or [], ensure_ascii=False)}"
    )
    lines.append(
        "  • SELECTED_UI_EVIDENCE_PT: "
        f"{json.dumps(pairing_diag.get('SELECTED_UI_EVIDENCE_PT') or [], ensure_ascii=False)}"
    )
    lines.append(
        "  • REJECTED_NON_BOM_UI_EVIDENCE: "
        f"{json.dumps(pairing_diag.get('REJECTED_NON_BOM_UI_EVIDENCE') or [], ensure_ascii=False)}"
    )
    lines.append(f"  • UE_SELECTION_REASON  : {pairing_diag.get('UE_SELECTION_REASON') or '-'}")
    lines.append(
        "  • FULL_CONFIGURATION_EVIDENCE: "
        f"{json.dumps(pairing_diag.get('FULL_CONFIGURATION_EVIDENCE') or [], ensure_ascii=False)}"
    )
    for key in (
        "FULL_CONFIGURATION_MATCHED",
        "FULL_CONFIGURATION_REQUESTED",
        "FULL_CONFIGURATION_CATALOG",
        "FULL_CONFIGURATION_SOURCE_FIELD",
        "FULL_CONFIGURATION_EVIDENCE_STRENGTH",
        "FULL_CONFIGURATION_PROVENANCE",
        "FULL_CONFIGURATION_PAGE",
        "FULL_CONFIGURATION_TABLE",
    ):
        value = pairing_diag.get(key)
        lines.append(f"  • {key}: {value if value is not None else '-'}")
    lines.append(
        "  • CONFIGURATION_CONFLICT_REASON: "
        f"{pairing_diag.get('CONFIGURATION_CONFLICT_REASON') or '-'}"
    )
    lines.append(
        "  • CONFIGURATION_CONFLICT_DETAILS: "
        f"{json.dumps(pairing_diag.get('CONFIGURATION_CONFLICT_DETAILS') or {}, ensure_ascii=False)}"
    )
    lines.append(f"  • MPN_FINAL_BLOCK_REASON: {pairing_diag.get('MPN_FINAL_BLOCK_REASON') or '-'}")

    compat_evidence = search_res.get("compatibility_evidence") or {}
    rel_type = compat_evidence.get("relation_type")
    lines.append(f"  • Relation Type        : {rel_type or 'Nessuna / Non applicabile'}")
    lines.append(f"  • Source / Provenance  : {compat_evidence.get('provenance') or '-'}")
    if compat_evidence.get("confidence") is not None:
        lines.append(f"  • Confidence           : {compat_evidence.get('confidence')}")

    connected_comps = compat_evidence.get("connected_components") or compat_evidence.get("componenti_collegati") or []
    if connected_comps:
        lines.append(f"  • Componenti Collegati ({len(connected_comps)}):")
        for cc in connected_comps:
            pt_c = cc.get("code") or "-"
            mfg_c = cc.get("mfg_code") or "-"
            role_c = cc.get("role") or "-"
            name_c = cc.get("name") or "-"
            note_c = f" [{cc.get('note')}]" if cc.get("note") else ""
            lines.append(f"      - PT: {pt_c} | MFG: {mfg_c} | Ruolo: {role_c} | Nome: {name_c}{note_c}")
    else:
        lines.append("  • Componenti Collegati : Nessun componente collegato")
    lines.append("")

    # 5. COMPONENT RELATIONS (separate sia dalla compatibilità UI/UE sia dalla BOM)
    lines.append(subsep)
    lines.append("5. INFORMAZIONI COMPONENTI / ACCESSORI")
    lines.append(subsep)
    component_relations = search_res.get("component_relations") or []
    component_products = search_res.get("component_relation_products") or []
    lines.append(
        f"  • Lookup Status         : "
        f"{search_res.get('component_relation_lookup_status') or 'NOT_APPLICABLE'}"
    )
    lines.append(
        f"  • Source                : "
        f"{search_res.get('component_relation_source') or 'catalog_component_relations_v3_4'}"
    )
    if component_products:
        lines.append(f"  • Product Context ({len(component_products)}):")
        for product in component_products:
            lines.append(
                "      - PT: {pt} | MFG: {mfg} | Ruolo: {role} | "
                "Table: {table} | Family: {family} | Nome: {name}".format(
                    pt=product.get("code") or "-",
                    mfg=product.get("mfg_code") or product.get("model") or "-",
                    role=product.get("role") or ("UI" if product.get("is_ui") else ("UE" if product.get("is_ue") else "-")),
                    table=product.get("table_id") or "-",
                    family=product.get("family_key") or "-",
                    name=product.get("name") or product.get("catalog_family") or "-",
                )
            )
    else:
        lines.append("  • Product Context       : Nessun prodotto clima identificato")

    if component_relations:
        lines.append(f"  • Component Relations ({len(component_relations)}):")
        for relation in component_relations:
            component = relation.get("component") or {}
            accessories = relation.get("accessories") or []
            if accessories:
                for accessory in accessories:
                    lines.append(
                        f"      - PT {accessory.get('pt') or '-'} | "
                        f"Modello: {accessory.get('model') or '-'}"
                    )
                    lines.append(f"        Tipo    : {component.get('type') or '-'}")
                    lines.append(
                        f"        Intent  : {component.get('relation_intent') or '-'}"
                    )
                    lines.append(
                        f"        Target  : {relation.get('attachment_target') or '-'}"
                    )
                    applies = relation.get("applies_to_models") or []
                    lines.append(
                        f"        Applies : {', '.join(str(model) for model in applies) if applies else '-'}"
                    )
                    lines.append(
                        f"        Lookup  : {relation.get('lookup_strategy') or '-'}"
                    )
                    lines.append("        Source  : catalog_component_relations_v3_4")
            else:
                lines.append(
                    f"      - Nessun PT separato | Tipo: {component.get('type') or '-'} | "
                    f"Intent: {component.get('relation_intent') or '-'} | "
                    f"Target: {relation.get('attachment_target') or '-'}"
                )
                if relation.get("raw_text"):
                    lines.append(f"        Nota    : {relation.get('raw_text')}")
    else:
        lines.append("  • Component Relations  : Nessuna relazione componente/accessorio")
    lines.append("")

    lines.append(sep)
    return "\n".join(lines)


def run_batch_qa(engine: CatalogSearchEngine, input_file: str, output_file: str):
    print(f"Avvio batch QA da: {input_file}")
    in_path = Path(input_file)
    if not in_path.exists():
        print(f"ERRORE: File di input '{input_file}' non trovato.")
        return

    rows_to_process = []
    is_xlsx = in_path.suffix.lower() in [".xlsx", ".xlsm"]

    if is_xlsx:
        import openpyxl
        wb = openpyxl.load_workbook(in_path, data_only=True)
        ws = wb.active
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(row):
                row_dict = {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
                rows_to_process.append(row_dict)
    else:
        with open(in_path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for r in reader:
                rows_to_process.append(r)

    print(f"Caricate {len(rows_to_process)} righe da analizzare.")

    out_records = []
    for idx, r in enumerate(rows_to_process, 1):
        if idx % 20 == 0 or idx == len(rows_to_process):
            print(f"  Elaborate {idx}/{len(rows_to_process)}...")

        # Individua campo titolo/query
        title = (
            r.get("titolo") or r.get("Titolo") or r.get("title") or r.get("query") or
            r.get("descrizione") or r.get("Descrizione") or r.get("nome") or list(r.values())[0]
        )
        title = str(title or "").strip()
        ref = r.get("riferimento") or r.get("codice_fornitore") or r.get("id") or r.get("ID")
        mpn = r.get("mpn") or r.get("mpn_attuale") or r.get("mfg_code")

        if not title:
            continue

        res = engine.search(title, limit=20)
        pools = res.get("candidate_pools", {})
        q_ctx = res.get("query_analysis", {}).get("query_context", {})

        # Estrai BOM candidata direttamente dal motore
        bom = res.get("bom", [])
        ue_item = next((it for it in bom if it.get("role") == "UE" or it.get("is_ue")), {})
        ui_items = [it for it in bom if it.get("role") == "UI" or it.get("is_ui")]

        compat_ev = res.get("compatibility_evidence") or {}

        rec = dict(r)
        rec["RETRIEVAL_BRAND"] = res.get("detected_brand")
        rec["RETRIEVAL_MATCH_TYPE"] = res.get("match_type")
        rec["PRODUCT_SCOPE"] = res.get("product_scope")
        rec["COMPATIBILITY_STATUS"] = res.get("compatibility_status")
        rec["COMPATIBILITY_RELATION"] = compat_ev.get("relation_type")
        rec["COMPATIBILITY_PROVENANCE"] = compat_ev.get("provenance")
        rec["PAIRING_REASON"] = res.get("pairing_reason")
        rec["PAIRING_STATUS"] = res.get("pairing_status")
        rec["PAIRING_SCORE_BREAKDOWN"] = json.dumps(
            res.get("pairing_score_breakdown") or {}, ensure_ascii=False
        )
        rec["MPN_FINAL_ALLOWED"] = bool(res.get("mpn_final_allowed"))
        pairing_diag = res.get("pairing_diagnostics") or {}
        rec["PRODUCT_IDENTITY_STATUS"] = pairing_diag.get("PRODUCT_IDENTITY_STATUS")
        rec["CONFIGURATION_STATUS"] = pairing_diag.get("CONFIGURATION_STATUS")
        rec["EXPLICIT_UE_TOKEN"] = pairing_diag.get("EXPLICIT_UE_TOKEN")
        rec["EXPLICIT_UE_MATCH_TYPE"] = pairing_diag.get("EXPLICIT_UE_MATCH_TYPE")
        rec["EXPLICIT_UE_CANDIDATES"] = json.dumps(
            pairing_diag.get("EXPLICIT_UE_CANDIDATES") or [], ensure_ascii=False
        )
        rec["SELECTED_UI_EVIDENCE_PT"] = json.dumps(
            pairing_diag.get("SELECTED_UI_EVIDENCE_PT") or [], ensure_ascii=False
        )
        rec["REJECTED_NON_BOM_UI_EVIDENCE"] = json.dumps(
            pairing_diag.get("REJECTED_NON_BOM_UI_EVIDENCE") or [], ensure_ascii=False
        )
        rec["UE_SELECTION_REASON"] = pairing_diag.get("UE_SELECTION_REASON")
        rec["FULL_CONFIGURATION_EVIDENCE"] = json.dumps(
            pairing_diag.get("FULL_CONFIGURATION_EVIDENCE") or [], ensure_ascii=False
        )
        for key in (
            "FULL_CONFIGURATION_MATCHED",
            "FULL_CONFIGURATION_REQUESTED",
            "FULL_CONFIGURATION_CATALOG",
            "FULL_CONFIGURATION_SOURCE_FIELD",
            "FULL_CONFIGURATION_EVIDENCE_STRENGTH",
            "FULL_CONFIGURATION_PROVENANCE",
            "FULL_CONFIGURATION_PAGE",
            "FULL_CONFIGURATION_TABLE",
        ):
            rec[key] = pairing_diag.get(key)
        rec["CONFIGURATION_CONFLICT_REASON"] = pairing_diag.get(
            "CONFIGURATION_CONFLICT_REASON"
        )
        rec["CONFIGURATION_CONFLICT_DETAILS"] = json.dumps(
            pairing_diag.get("CONFIGURATION_CONFLICT_DETAILS") or {}, ensure_ascii=False
        )
        rec["MPN_FINAL_BLOCK_REASON"] = pairing_diag.get("MPN_FINAL_BLOCK_REASON")
        rec["COMPONENT_RELATION_LOOKUP_STATUS"] = res.get("component_relation_lookup_status")
        rec["COMPONENT_RELATION_SOURCE"] = res.get("component_relation_source")
        rec["COMPONENT_RELATIONS_JSON"] = json.dumps(
            res.get("component_relations") or [], ensure_ascii=False
        )
        rec["CANDIDATE_UE_PT"] = ue_item.get("code")
        rec["CANDIDATE_UE_MFG"] = ue_item.get("mfg_code")
        rec["CANDIDATE_UE_NAME"] = ue_item.get("name")
        rec["CANDIDATE_UE_TIER"] = ue_item.get("evidence_tier")
        rec["QUERY_COLOR"] = q_ctx.get("query_color")
        rec["CANDIDATE_UE_COLOR"] = ue_item.get("candidate_color")
        rec["CANDIDATE_UE_VARIANT_CONFLICT"] = bool(ue_item.get("variant_conflict"))

        for u_idx, ui_it in enumerate(ui_items, 1):
            rec[f"CANDIDATE_UI{u_idx}_PT"] = ui_it.get("code")
            rec[f"CANDIDATE_UI{u_idx}_MFG"] = ui_it.get("mfg_code")
            rec[f"CANDIDATE_UI{u_idx}_NAME"] = ui_it.get("name")
            rec[f"CANDIDATE_UI{u_idx}_BTU"] = ui_it.get("taglia_btu")
            rec[f"CANDIDATE_UI{u_idx}_TIER"] = ui_it.get("evidence_tier")
            rec[f"CANDIDATE_UI{u_idx}_COLOR"] = ui_it.get("candidate_color")
            rec[f"CANDIDATE_UI{u_idx}_VARIANT_CONFLICT"] = bool(
                ui_it.get("variant_conflict")
            )

        out_records.append(rec)

    # Scrittura output
    out_p = Path(output_file)
    if out_p.suffix.lower() == ".xlsx":
        import openpyxl
        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = "QA_Results"

        if out_records:
            fieldnames = list(out_records[0].keys())
            ws_out.append(fieldnames)
            for rec in out_records:
                ws_out.append([rec.get(k) for k in fieldnames])
        wb_out.save(out_p)
    else:
        with open(out_p, "w", newline="", encoding="utf-8-sig") as f:
            if out_records:
                writer = csv.DictWriter(f, fieldnames=list(out_records[0].keys()))
                writer.writeheader()
                writer.writerows(out_records)

    print(f"Batch QA completato con successo. Risultati salvati in: {output_file}")


def main():
    parser = argparse.ArgumentParser(description="Strumento Manual QA per test e verifica del Candidate Retrieval.")
    parser.add_argument("--title", type=str, help="Titolo della query per test singolo")
    parser.add_argument("--ref", type=str, help="Riferimento opzionale (es. codice riga o codice fornitore)")
    parser.add_argument("--mpn", type=str, help="MPN attuale opzionale")
    parser.add_argument("--batch", type=str, help="Percorso del file input XLSX o CSV per test batch")
    parser.add_argument("--output", type=str, default="qa_results.xlsx", help="Percorso file output per test batch")
    parser.add_argument("--interactive", action="store_true", help="Avvia modalità interattiva da riga di comando")

    args = parser.parse_args()

    print("Inizializzazione CatalogSearchEngine per QA...")
    engine = CatalogSearchEngine()
    engine._ensure_initialized()

    if args.batch:
        run_batch_qa(engine, args.batch, args.output)
    elif args.title:
        res = engine.search(args.title, limit=20)
        report = format_qa_single_report(args.title, res, reference=args.ref, current_mpn=args.mpn)
        print("\n" + report)
    elif args.interactive or len(sys.argv) == 1:
        print("\n=== MODALITA' INTERATTIVA MANUAL QA (Digita 'exit' per uscire) ===")
        while True:
            try:
                title = input("\nInserisci Titolo Query: ").strip()
                if not title or title.lower() in ["exit", "quit"]:
                    break
                ref = input("Inserisci Riferimento (opzionale, premi INVIO per saltare): ").strip() or None
                mpn = input("Inserisci MPN attuale (opzionale, premi INVIO per saltare): ").strip() or None

                res = engine.search(title, limit=20)
                report = format_qa_single_report(title, res, reference=ref, current_mpn=mpn)
                print("\n" + report)
            except (KeyboardInterrupt, EOFError):
                break
        print("\nSessione QA terminata.")
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
